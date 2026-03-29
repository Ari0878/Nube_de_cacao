# services/restore_service.py
import os
import json
from datetime import datetime
from io import BytesIO
import openpyxl
from db import db
from bson import ObjectId

# Directorio para almacenar archivos de restauración temporales
RESTORE_DIR = "restore_temp"
if not os.path.exists(RESTORE_DIR):
    os.makedirs(RESTORE_DIR)

HISTORIAL_RESTORE_FILE = os.path.join(RESTORE_DIR, "historial_restauraciones.json")


def cargar_historial_restauraciones():
    """Carga el historial de restauraciones"""
    if os.path.exists(HISTORIAL_RESTORE_FILE):
        with open(HISTORIAL_RESTORE_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []


def guardar_historial_restauraciones(historial):
    """Guarda el historial de restauraciones"""
    with open(HISTORIAL_RESTORE_FILE, 'w', encoding='utf-8') as f:
        json.dump(historial, f, ensure_ascii=False, indent=2)


def agregar_historial_restauracion(tipo_archivo, colecciones_restauradas, registros_totales, estado="exitoso", error=None):
    """Agrega una entrada al historial de restauraciones"""
    historial = cargar_historial_restauraciones()
    
    entrada = {
        "fecha": datetime.now().isoformat(),
        "tipo_archivo": tipo_archivo,
        "colecciones": colecciones_restauradas,
        "registros_totales": registros_totales,
        "estado": estado,
        "error": error
    }
    
    historial.insert(0, entrada)
    
    # Mantener solo los últimos 50 registros
    if len(historial) > 50:
        historial = historial[:50]
    
    guardar_historial_restauraciones(historial)


def validar_archivo_restauracion(filename):
    """
    Valida que el archivo tenga una extensión permitida
    Retorna: (valido: bool, extension: str, mensaje: str)
    """
    extensiones_permitidas = {'.sql', '.xlsx', '.json'}
    
    if not filename:
        return False, None, "No se proporcionó ningún archivo"
    
    extension = os.path.splitext(filename)[1].lower()
    
    if extension not in extensiones_permitidas:
        return False, extension, f"Extensión no permitida. Solo se aceptan: {', '.join(extensiones_permitidas)}"
    
    return True, extension, "Archivo válido"


def restaurar_desde_sql(file_content):
    """
    Restaura la base de datos desde un archivo SQL
    Retorna: (success: bool, mensaje: str, stats: dict)
    """
    try:
        # Decodificar contenido
        sql_content = file_content.decode('utf-8')
        
        # Parsear el SQL y extraer datos JSON
        colecciones_restauradas = {}
        registros_totales = 0
        
        # Dividir en líneas y procesar
        lineas = sql_content.split('\n')
        tabla_actual = None
        
        for linea in lineas:
            linea = linea.strip()
            
            # Detectar tabla actual
            if linea.startswith('CREATE TABLE IF NOT EXISTS'):
                # Extraer nombre de tabla
                partes = linea.split('`')
                if len(partes) >= 2:
                    tabla_actual = partes[1]
                    colecciones_restauradas[tabla_actual] = []
            
            # Procesar INSERT
            elif linea.startswith('INSERT INTO') and tabla_actual:
                try:
                    # Extraer el JSON del INSERT
                    # Formato: INSERT INTO `tabla` (`mongo_id`, `datos`) VALUES ('id', 'json');
                    inicio_json = linea.find("', '") + 4
                    fin_json = linea.rfind("');")
                    
                    if inicio_json > 3 and fin_json > inicio_json:
                        json_str = linea[inicio_json:fin_json]
                        # Revertir escape de comillas
                        json_str = json_str.replace("''", "'")
                        
                        # Parsear JSON
                        documento = json.loads(json_str)
                        colecciones_restauradas[tabla_actual].append(documento)
                        registros_totales += 1
                
                except Exception as e:
                    print(f"Error al procesar línea SQL: {e}")
                    continue
        
        # Insertar en MongoDB
        if db is None:
            return False, "No hay conexión a la base de datos", {}
        
        for nombre_coleccion, documentos in colecciones_restauradas.items():
            if not documentos:
                continue
            
            try:
                coleccion = db[nombre_coleccion]
                
                # Limpiar colección existente (opcional - comentar si no se desea)
                # coleccion.delete_many({})
                
                # Convertir _id string a ObjectId
                for doc in documentos:
                    if '_id' in doc and isinstance(doc['_id'], str):
                        try:
                            doc['_id'] = ObjectId(doc['_id'])
                        except:
                            # Si no es un ObjectId válido, dejarlo como string
                            pass
                
                # Insertar documentos
                coleccion.insert_many(documentos, ordered=False)
            
            except Exception as e:
                print(f"Error al restaurar colección {nombre_coleccion}: {e}")
                continue
        
        stats = {
            "colecciones": len(colecciones_restauradas),
            "registros": registros_totales,
            "detalles": {nombre: len(docs) for nombre, docs in colecciones_restauradas.items()}
        }
        
        agregar_historial_restauracion("SQL", list(colecciones_restauradas.keys()), registros_totales)
        
        return True, f"Restauración exitosa: {registros_totales} registros en {len(colecciones_restauradas)} colecciones", stats
    
    except Exception as e:
        error_msg = f"Error al restaurar desde SQL: {str(e)}"
        print(error_msg)
        import traceback
        traceback.print_exc()
        agregar_historial_restauracion("SQL", [], 0, "error", error_msg)
        return False, error_msg, {}


def restaurar_desde_excel(file_content):
    """
    Restaura la base de datos desde un archivo Excel
    Retorna: (success: bool, mensaje: str, stats: dict)
    """
    try:
        # Cargar Excel
        wb = openpyxl.load_workbook(BytesIO(file_content))
        
        colecciones_restauradas = {}
        registros_totales = 0
        
        # Procesar cada hoja (excepto la de información)
        for sheet_name in wb.sheetnames:
            if sheet_name.startswith('📊'):  # Saltar hoja de información
                continue
            
            ws = wb[sheet_name]
            
            # Obtener encabezados (primera fila)
            encabezados = []
            for cell in ws[1]:
                if cell.value:
                    encabezados.append(cell.value)
            
            if not encabezados:
                continue
            
            # Procesar datos
            documentos = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):  # Saltar filas vacías
                    continue
                
                documento = {}
                for idx, valor in enumerate(row):
                    if idx < len(encabezados):
                        clave = encabezados[idx]
                        
                        # Intentar parsear JSON si parece serlo
                        if isinstance(valor, str) and (valor.startswith('{') or valor.startswith('[')):
                            try:
                                valor = json.loads(valor)
                            except:
                                pass
                        
                        documento[clave] = valor
                
                if documento:
                    documentos.append(documento)
                    registros_totales += 1
            
            colecciones_restauradas[sheet_name] = documentos
        
        # Insertar en MongoDB
        if db is None:
            return False, "No hay conexión a la base de datos", {}
        
        for nombre_coleccion, documentos in colecciones_restauradas.items():
            if not documentos:
                continue
            
            try:
                coleccion = db[nombre_coleccion]
                
                # Convertir _id a ObjectId si es necesario
                for doc in documentos:
                    if '_id' in doc and isinstance(doc['_id'], str):
                        try:
                            doc['_id'] = ObjectId(doc['_id'])
                        except:
                            pass
                
                # Insertar documentos
                coleccion.insert_many(documentos, ordered=False)
            
            except Exception as e:
                print(f"Error al restaurar colección {nombre_coleccion}: {e}")
                continue
        
        stats = {
            "colecciones": len(colecciones_restauradas),
            "registros": registros_totales,
            "detalles": {nombre: len(docs) for nombre, docs in colecciones_restauradas.items()}
        }
        
        agregar_historial_restauracion("EXCEL", list(colecciones_restauradas.keys()), registros_totales)
        
        return True, f"Restauración exitosa: {registros_totales} registros en {len(colecciones_restauradas)} colecciones", stats
    
    except Exception as e:
        error_msg = f"Error al restaurar desde Excel: {str(e)}"
        print(error_msg)
        import traceback
        traceback.print_exc()
        agregar_historial_restauracion("EXCEL", [], 0, "error", error_msg)
        return False, error_msg, {}


def restaurar_desde_json(file_content):
    """
    Restaura la base de datos desde un archivo JSON
    Retorna: (success: bool, mensaje: str, stats: dict)
    """
    try:
        # Decodificar y parsear JSON
        json_str = file_content.decode('utf-8')
        data = json.loads(json_str)
        
        # Verificar estructura
        if 'colecciones' not in data:
            return False, "El archivo JSON no tiene la estructura correcta (falta 'colecciones')", {}
        
        colecciones = data['colecciones']
        registros_totales = 0
        
        # Insertar en MongoDB
        if db is None:
            return False, "No hay conexión a la base de datos", {}
        
        for nombre_coleccion, documentos in colecciones.items():
            if not documentos:
                continue
            
            try:
                coleccion = db[nombre_coleccion]
                
                # Convertir _id a ObjectId
                for doc in documentos:
                    if '_id' in doc and isinstance(doc['_id'], str):
                        try:
                            doc['_id'] = ObjectId(doc['_id'])
                        except:
                            pass
                
                # Insertar documentos
                coleccion.insert_many(documentos, ordered=False)
                registros_totales += len(documentos)
            
            except Exception as e:
                print(f"Error al restaurar colección {nombre_coleccion}: {e}")
                continue
        
        stats = {
            "colecciones": len(colecciones),
            "registros": registros_totales,
            "detalles": {nombre: len(docs) for nombre, docs in colecciones.items()}
        }
        
        agregar_historial_restauracion("JSON", list(colecciones.keys()), registros_totales)
        
        return True, f"Restauración exitosa: {registros_totales} registros en {len(colecciones)} colecciones", stats
    
    except Exception as e:
        error_msg = f"Error al restaurar desde JSON: {str(e)}"
        print(error_msg)
        import traceback
        traceback.print_exc()
        agregar_historial_restauracion("JSON", [], 0, "error", error_msg)
        return False, error_msg, {}


def obtener_estadisticas_restauraciones():
    """Obtiene estadísticas del historial de restauraciones"""
    historial = cargar_historial_restauraciones()
    
    if not historial:
        return {
            "total_restauraciones": 0,
            "exitosas": 0,
            "fallidas": 0,
            "ultima_restauracion": None,
            "total_registros_restaurados": 0
        }
    
    exitosas = sum(1 for h in historial if h.get("estado") == "exitoso")
    fallidas = len(historial) - exitosas
    total_registros = sum(h.get("registros_totales", 0) for h in historial if h.get("estado") == "exitoso")
    
    return {
        "total_restauraciones": len(historial),
        "exitosas": exitosas,
        "fallidas": fallidas,
        "ultima_restauracion": historial[0].get("fecha") if historial else None,
        "total_registros_restaurados": total_registros
    }